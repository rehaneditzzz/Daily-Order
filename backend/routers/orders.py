import io
import uuid
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pymongo import ReturnDocument

from lib.db import db
from models.order import (
    Activity,
    DashboardStats,
    Item,
    ItemIn,
    ItemUpdate,
    MarkAvailable,
    Order,
    OrderCreate,
    OrderUpdate,
    ProcurementEntry,
    ProcurementGroup,
    Shipment,
    ShipmentCreate,
    ShipmentUpdate,
    now_iso,
)

router = APIRouter()

UPLOAD_DIR = Path(__file__).parent.parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

STAGE_LABELS = {
    "on_hold": "On Hold",
    "payment_pending": "Payment Pending",
    "procurement": "Procurement",
    "billing": "Billing",
    "packing": "Packing",
    "tracking": "Tracking",
    "shipped": "Shipped",
    "delivered": "Delivered",
}


def compute_stage(o: Order) -> str:
    if o.on_hold:
        return "on_hold"
    if not o.payment_received:
        return "payment_pending"
    if any(i.availability != "available" for i in o.items):
        return "procurement"
    if not o.shipments or not all(s.bill_ready for s in o.shipments):
        return "billing"
    if not all(s.packed for s in o.shipments):
        return "packing"
    if not all(s.tracking_added for s in o.shipments):
        return "tracking"
    if all(s.delivered for s in o.shipments):
        return "delivered"
    return "shipped"


def out(o: Order) -> Order:
    o.stage = compute_stage(o)
    return o


async def next_order_no() -> int:
    doc = await db.counters.find_one_and_update(
        {"_id": "order_no"}, {"$inc": {"value": 1}}, upsert=True, return_document=ReturnDocument.AFTER
    )
    return 1000 + int(doc["value"])


async def load(order_id: str) -> Order:
    doc = await db.orders.find_one({"id": order_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Order not found")
    return Order(**doc)


async def save(o: Order, actor: str, text: str) -> Order:
    o.activity.insert(0, Activity(actor=actor or "Team", text=text))
    o.activity = o.activity[:300]
    o.updated_at = now_iso()
    o.stage = compute_stage(o)
    await db.orders.replace_one({"id": o.id}, o.model_dump(), upsert=True)
    return o


def find_item(o: Order, item_id: str) -> Item:
    for i in o.items:
        if i.id == item_id:
            return i
    raise HTTPException(status_code=404, detail="Item not found")


def find_shipment(o: Order, sid: str) -> Shipment:
    for s in o.shipments:
        if s.id == sid:
            return s
    raise HTTPException(status_code=404, detail="Shipment not found")


def fmt_qty(q: float) -> str:
    return str(int(q)) if float(q).is_integer() else str(q)


# ---------------- Orders ----------------
@router.get("/orders", response_model=list[Order])
async def list_orders(
    q: str = "",
    stage: str = "",
    assigned: str = "",
    country: str = "",
    speed: str = "",
):
    docs = await db.orders.find({}, {"_id": 0}).sort("updated_at", -1).to_list(5000)
    orders = [out(Order(**d)) for d in docs]
    needle = q.strip().lower()
    res = []
    for o in orders:
        if stage == "open" and o.stage == "delivered":
            continue
        if stage and stage != "open" and o.stage != stage:
            continue
        if assigned and o.assigned_to != assigned:
            continue
        if country and o.country.lower() != country.lower():
            continue
        if speed and not any(s.speed == speed for s in o.shipments):
            continue
        if needle:
            hay = " ".join(
                [str(o.order_no), o.client_name, o.country, o.client_phone, o.client_email, o.notes]
                + [i.name for i in o.items]
                + [s.address + " " + s.recipient + " " + " ".join(s.tracking_ids) for s in o.shipments]
            ).lower()
            if needle not in hay:
                continue
        res.append(o)
    return res


@router.post("/orders", response_model=Order)
async def create_order(body: OrderCreate):
    items = [Item(**i.model_dump(exclude={"actor"})) for i in body.items]
    n = body.shipment_count
    shipments = []
    for k in range(n):
        alloc: dict[str, float] = {}
        for it in items:
            if float(it.qty).is_integer():
                base, rem = divmod(int(it.qty), n)
                share = base + (1 if k < rem else 0)
            else:
                share = it.qty if k == 0 else 0
            if share:
                alloc[it.id] = share
        shipments.append(
            Shipment(
                label=f"Parcel {k + 1}",
                recipient=body.client_name,
                address=body.address,
                country=body.country,
                speed=body.speed,
                allocations=alloc,
            )
        )
    data = body.model_dump(exclude={"items", "address", "shipment_count", "speed", "actor"})
    o = Order(order_no=await next_order_no(), items=items, shipments=shipments, **data)
    return await save(o, body.actor, f"Created order with {len(items)} item(s) and {n} parcel(s)")


@router.get("/orders/{order_id}", response_model=Order)
async def get_order(order_id: str):
    return out(await load(order_id))


@router.patch("/orders/{order_id}", response_model=Order)
async def update_order(order_id: str, body: OrderUpdate):
    o = await load(order_id)
    changes = body.model_dump(exclude_unset=True, exclude={"actor"})
    if not changes:
        return out(o)
    for k, v in changes.items():
        setattr(o, k, v)
    if "payment_received" in changes:
        text = "Marked payment as received" if body.payment_received else "Marked payment as pending"
    elif "on_hold" in changes:
        text = f"Put order on hold: {o.hold_reason}" if body.on_hold else "Released order from hold"
    elif set(changes) == {"show_hts"}:
        text = "Showed HTS codes" if body.show_hts else "Hid HTS codes"
    elif "assigned_to" in changes and len(changes) == 1:
        text = f"Assigned order to {body.assigned_to or 'nobody'}"
    else:
        text = "Updated " + ", ".join(k.replace("_", " ") for k in changes)
    return await save(o, body.actor, text)


@router.delete("/orders/{order_id}")
async def delete_order(order_id: str):
    res = await db.orders.delete_one({"id": order_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Order not found")
    return {"ok": True}


# ---------------- Items ----------------
@router.post("/orders/{order_id}/items", response_model=Order)
async def add_item(order_id: str, body: ItemIn):
    o = await load(order_id)
    item = Item(**body.model_dump(exclude={"actor"}))
    o.items.append(item)
    return await save(o, body.actor, f"Added item {item.name} — {fmt_qty(item.qty)} {item.unit}".strip())


@router.patch("/orders/{order_id}/items/{item_id}", response_model=Order)
async def update_item(order_id: str, item_id: str, body: ItemUpdate):
    o = await load(order_id)
    item = find_item(o, item_id)
    changes = body.model_dump(exclude_unset=True, exclude={"actor"})
    for k, v in changes.items():
        setattr(item, k, v)
    if set(changes) == {"availability"}:
        label = {"available": "Available", "need_market": "Need from market", "unavailable": "Unavailable"}
        text = f"{item.name}: marked {label[item.availability]}"
    else:
        text = f"Edited item {item.name}"
    return await save(o, body.actor, text)


@router.delete("/orders/{order_id}/items/{item_id}", response_model=Order)
async def delete_item(order_id: str, item_id: str, actor: str = ""):
    o = await load(order_id)
    item = find_item(o, item_id)
    o.items = [i for i in o.items if i.id != item_id]
    for s in o.shipments:
        s.allocations.pop(item_id, None)
    return await save(o, actor, f"Removed item {item.name}")


# ---------------- Shipments ----------------
@router.post("/orders/{order_id}/shipments", response_model=Order)
async def add_shipment(order_id: str, body: ShipmentCreate):
    o = await load(order_id)
    s = Shipment(
        label=f"Parcel {len(o.shipments) + 1}",
        recipient=body.recipient or o.client_name,
        address=body.address,
        country=body.country or o.country,
        speed=body.speed,
    )
    o.shipments.append(s)
    return await save(o, body.actor, f"Added {s.label}")


@router.patch("/orders/{order_id}/shipments/{sid}", response_model=Order)
async def update_shipment(order_id: str, sid: str, body: ShipmentUpdate):
    o = await load(order_id)
    s = find_shipment(o, sid)
    changes = body.model_dump(exclude_unset=True, exclude={"actor"})
    if "tracking_ids" in changes:
        changes["tracking_ids"] = [t.strip() for t in changes["tracking_ids"] if t.strip()]
        if changes["tracking_ids"] and "tracking_added" not in changes:
            changes["tracking_added"] = True
    if "allocations" in changes:
        changes["allocations"] = {k: v for k, v in changes["allocations"].items() if v and v > 0}
    for k, v in changes.items():
        setattr(s, k, v)
    flags = {"bill_ready": "Bill ready", "packed": "Packed", "tracking_added": "Tracking added", "delivered": "Delivered"}
    if len(changes) == 1 and next(iter(changes)) in flags:
        key = next(iter(changes))
        text = f"{s.label}: {flags[key]} {'✓' if changes[key] else 'unchecked'}"
    elif "tracking_ids" in changes:
        text = f"{s.label}: tracking updated ({', '.join(s.tracking_ids) or 'none'})"
    elif "allocations" in changes:
        text = f"{s.label}: updated item split"
    else:
        text = f"{s.label}: updated " + ", ".join(k.replace("_", " ") for k in changes)
    return await save(o, body.actor, text)


@router.delete("/orders/{order_id}/shipments/{sid}", response_model=Order)
async def delete_shipment(order_id: str, sid: str, actor: str = ""):
    o = await load(order_id)
    s = find_shipment(o, sid)
    o.shipments = [x for x in o.shipments if x.id != sid]
    return await save(o, actor, f"Removed {s.label}")


@router.post("/orders/{order_id}/shipments/{sid}/photos", response_model=Order)
async def upload_photo(order_id: str, sid: str, file: UploadFile = File(...), actor: str = Form("")):
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(status_code=400, detail="Only image files are allowed")
    o = await load(order_id)
    s = find_shipment(o, sid)
    ext = Path(file.filename or "photo.jpg").suffix.lower() or ".jpg"
    name = f"{uuid.uuid4().hex}{ext}"
    data = await file.read()
    if len(data) > 8 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Image too large (max 8MB)")
    # Stored in MongoDB (not disk) so photos survive restarts on free hosting.
    await db.photos.insert_one({"name": name, "content_type": file.content_type, "data": data, "created_at": now_iso()})
    s.photos.append(f"/api/uploads/{name}")
    return await save(o, actor, f"{s.label}: uploaded parcel photo")


@router.delete("/orders/{order_id}/shipments/{sid}/photos", response_model=Order)
async def delete_photo(order_id: str, sid: str, url: str = Query(...), actor: str = ""):
    o = await load(order_id)
    s = find_shipment(o, sid)
    if url not in s.photos:
        raise HTTPException(status_code=404, detail="Photo not found")
    s.photos = [p for p in s.photos if p != url]
    await db.photos.delete_one({"name": url.rsplit("/", 1)[-1]})
    return await save(o, actor, f"{s.label}: removed a photo")


@router.get("/uploads/{name}")
async def get_upload(name: str):
    name = Path(name).name
    doc = await db.photos.find_one({"name": name})
    if doc:
        return Response(
            content=bytes(doc["data"]),
            media_type=doc.get("content_type") or "image/jpeg",
            headers={"Cache-Control": "public, max-age=31536000, immutable"},
        )
    path = UPLOAD_DIR / name  # legacy files saved to disk before the DB switch
    if not path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(path)


# ---------------- Dashboard / Procurement ----------------
@router.get("/dashboard", response_model=DashboardStats)
async def dashboard():
    docs = await db.orders.find({}, {"_id": 0}).sort("updated_at", -1).to_list(5000)
    orders = [out(Order(**d)) for d in docs]
    by_stage = {k: 0 for k in STAGE_LABELS}
    by_country: dict[str, int] = {}
    pending_items = bill = pack = track = express = 0
    open_orders = 0
    for o in orders:
        by_stage[o.stage] = by_stage.get(o.stage, 0) + 1
        if o.stage == "delivered":
            continue
        open_orders += 1
        c = o.country or "Unknown"
        by_country[c] = by_country.get(c, 0) + 1
        pending_items += sum(1 for i in o.items if i.availability != "available")
        for s in o.shipments:
            if s.delivered:
                continue
            bill += 0 if s.bill_ready else 1
            pack += 0 if s.packed else 1
            track += 1 if (s.packed and not s.tracking_added) else 0
            express += 1 if s.speed == "express" else 0
    return DashboardStats(
        total_orders=len(orders),
        open_orders=open_orders,
        by_stage=by_stage,
        pending_items=pending_items,
        bill_pending=bill,
        packing_due=pack,
        tracking_needed=track,
        express_open=express,
        by_country=by_country,
        recent=orders[:6],
    )


@router.get("/procurement", response_model=list[ProcurementGroup])
async def procurement():
    docs = await db.orders.find({}, {"_id": 0}).to_list(5000)
    groups: dict[str, ProcurementGroup] = {}
    for d in docs:
        o = Order(**d)
        for i in o.items:
            if i.availability == "available":
                continue
            key = i.name.strip().lower()
            g = groups.get(key)
            if not g:
                g = ProcurementGroup(name=i.name.strip(), unit=i.unit, total_qty=0, need_market=0, unavailable=0, entries=[])
                groups[key] = g
            g.total_qty += i.qty
            if i.availability == "need_market":
                g.need_market += 1
            else:
                g.unavailable += 1
            g.entries.append(
                ProcurementEntry(
                    order_id=o.id, order_no=o.order_no, client_name=o.client_name, assigned_to=o.assigned_to,
                    item_id=i.id, qty=i.qty, unit=i.unit, availability=i.availability, note=i.note,
                )
            )
    return sorted(groups.values(), key=lambda g: (-len(g.entries), g.name))


@router.post("/procurement/mark-available", response_model=list[ProcurementGroup])
async def mark_available(body: MarkAvailable):
    key = body.name.strip().lower()
    docs = await db.orders.find({}, {"_id": 0}).to_list(5000)
    for d in docs:
        o = Order(**d)
        hit = False
        for i in o.items:
            if i.name.strip().lower() == key and i.availability != "available":
                i.availability = "available"
                hit = True
        if hit:
            await save(o, body.actor, f"{body.name}: procured from market (marked Available)")
    return await procurement()


# ---------------- Excel export ----------------
@router.get("/export/excel")
async def export_excel():
    docs = await db.orders.find({}, {"_id": 0}).sort("order_no", 1).to_list(10000)
    orders = [out(Order(**d)) for d in docs]
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")

    def sheet(ws, headers, rows):
        ws.append(headers)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 60)
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Orders"
    sheet(ws, ["Order #", "Client", "Phone", "Email", "Country", "Assigned To", "Stage", "Payment Received",
               "Amount", "Currency", "Payment Method", "On Hold", "Hold Reason", "Items", "Parcels", "Notes",
               "Created", "Updated"],
          [[o.order_no, o.client_name, o.client_phone, o.client_email, o.country, o.assigned_to,
            STAGE_LABELS.get(o.stage, o.stage), "Yes" if o.payment_received else "No", o.payment_amount,
            o.currency, o.payment_method, "Yes" if o.on_hold else "No", o.hold_reason, len(o.items),
            len(o.shipments), o.notes, o.created_at, o.updated_at] for o in orders])
    sheet(wb.create_sheet("Items"),
          ["Order #", "Client", "Medicine", "Qty", "Unit", "Price", "HTS Code", "Availability", "Note"],
          [[o.order_no, o.client_name, i.name, i.qty, i.unit, i.price, i.hts_code, i.availability, i.note]
           for o in orders for i in o.items])
    names = {}
    for o in orders:
        for i in o.items:
            names[i.id] = f"{i.name} x{fmt_qty(i.qty)}"
    sheet(wb.create_sheet("Parcels"),
          ["Order #", "Client", "Parcel", "Recipient", "Address", "Country", "Speed", "Courier", "Tracking IDs",
           "Tracking Notes", "Items in Parcel", "Bill Ready", "Packed", "Tracking Added", "Delivered", "Photos"],
          [[o.order_no, o.client_name, s.label, s.recipient, s.address, s.country, s.speed, s.courier,
            ", ".join(s.tracking_ids), s.tracking_notes,
            "; ".join(f"{names.get(k, k).split(' x')[0]} x{fmt_qty(v)}" for k, v in s.allocations.items()),
            "Yes" if s.bill_ready else "No", "Yes" if s.packed else "No",
            "Yes" if s.tracking_added else "No", "Yes" if s.delivered else "No", len(s.photos)]
           for o in orders for s in o.shipments])
    sheet(wb.create_sheet("Activity"), ["Order #", "Client", "When (UTC)", "By", "What"],
          [[o.order_no, o.client_name, a.at, a.actor, a.text] for o in orders for a in o.activity])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"orders-backup-{now_iso()[:10]}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
